#!/usr/bin/env python3
"""Extract text from an Excel workbook (.xlsx) received on stdin; emit JSON on stdout.

Mirrors extract_pdf.py's output shape so the route's search/pagination logic is
shared: each sheet is a "page" whose text is pipe-separated rows — the same
" | " column convention formulas parse with splitCols().

Requires openpyxl (pip install openpyxl).
"""
import datetime
import io
import json
import sys

# Bound the text volume: a million-row sheet must not produce a gigabyte of JSON.
MAX_ROWS_PER_SHEET = 5000


import re


def _cell(v):
    if v is None:
        return ''
    # Date-typed cells come back as midnight datetimes — render the date alone.
    if isinstance(v, datetime.datetime) and v.time() == datetime.time(0, 0):
        return v.date().isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    # One cell must stay one column of one line: newlines inside a cell would
    # split the row, a literal '|' would shift columns; ▲/▼ are formatting glyphs.
    return re.sub(r'\s+', ' ', str(v).replace('|', ' ').replace('▲', '').replace('▼', '')).strip()


def extract(data: bytes) -> dict:
    from openpyxl import load_workbook
    # Not read_only: we need ws.merged_cells to de-duplicate merged banners.
    # (openpyxl cannot render number formats, so values stay raw here — the
    # in-browser SheetJS extractor is the display-faithful primary path.)
    wb = load_workbook(io.BytesIO(data), data_only=True)
    pages = []
    names = []
    for i, ws in enumerate(wb.worksheets, 1):
        names.append(ws.title)
        # Merged regions: exports often duplicate the value into every covered
        # cell — blank everything but the anchor so a banner appears once.
        blanked = set()
        for m in ws.merged_cells.ranges:
            for r in range(m.min_row, m.max_row + 1):
                for c in range(m.min_col, m.max_col + 1):
                    if r == m.min_row and c == m.min_col:
                        continue
                    blanked.add((r, c))
        lines = []
        truncated = False
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            if r > MAX_ROWS_PER_SHEET:
                truncated = True
                break
            cells = [
                '' if (r, c) in blanked else _cell(v)
                for c, v in enumerate(row, 1)
            ]
            while cells and cells[-1] == '':
                cells.pop()
            if not cells:
                continue
            lines.append(' | '.join(cells))
        if truncated:
            lines.append(f'[... sheet truncated after {MAX_ROWS_PER_SHEET} rows ...]')
        pages.append({'pageNum': i, 'label': ws.title, 'text': '\n'.join(lines)})
    wb.close()
    return {
        'pages': pages,
        'metadata': {
            'totalSheets': len(pages),
            'sheetNames': names,
        },
    }


if __name__ == '__main__':
    data = sys.stdin.buffer.read()
    if not data:
        json.dump({'error': 'No data received on stdin'}, sys.stdout)
        sys.exit(1)
    try:
        result = extract(data)
    except ImportError:
        json.dump({'error': 'openpyxl is not installed on the proxy — pip install openpyxl'}, sys.stdout)
        sys.exit(1)
    except Exception as e:  # noqa: BLE001 — surface any parse failure as JSON
        json.dump({'error': f'xlsx extraction failed: {e}'}, sys.stdout)
        sys.exit(1)
    json.dump(result, sys.stdout, ensure_ascii=False)
